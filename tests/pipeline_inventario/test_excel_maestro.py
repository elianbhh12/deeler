import csv
from pathlib import Path

from openpyxl import load_workbook

from python_pipeline import excel_report


def _write_group_csv(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(["Nombre de los JSONs", "Nombre TA config", "Nombre del subtipo",
                          "Cantidad de archivos con ese subtipo", "Proceso", "Ruta JSON R3",
                          "Observaciones", "Transmisiones", "Tipo"])
        writer.writerow(["AAA.json", "ta1", "carta", "1", "p1", "s3://x/AAA", "", "Si", "R3 - Topics"])
        writer.writerow([])
        writer.writerow(["TOTAL_R3", "", "", "1", "", "", "", "", ""])


def test_build_maestro_workbook(tmp_path):
    group_csv = tmp_path / "grupo.csv"
    _write_group_csv(group_csv)

    out = tmp_path / "maestro.xlsx"
    excel_report.build_maestro_workbook(out, group_csv, resumen={"Ambiente": "qa", "Total de flujos conocidos (presentes)": 1})

    wb = load_workbook(out)
    assert wb.sheetnames == ["Resumen", "Reporte Agrupado"]

    ws_resumen = wb["Resumen"]
    valores = {row[0]: row[1] for row in ws_resumen.iter_rows(min_row=2, values_only=True)}
    assert valores["Ambiente"] == "qa"

    ws_grupo = wb["Reporte Agrupado"]
    rows = list(ws_grupo.iter_rows(values_only=True))
    assert rows[1][0] == "AAA.json"
    assert rows[1][2] == "carta"


def test_build_maestro_workbook_con_r2_agrega_hoja_propia(tmp_path):
    """Los R2 acumulados (ver pipeline_service.py) quedan en su propia hoja
    del maestro, no se pierden entre corridas como pasaba antes (solo vivían
    en el reporte_completo.xlsx de esa corrida puntual)."""
    group_csv = tmp_path / "grupo.csv"
    _write_group_csv(group_csv)

    out = tmp_path / "maestro.xlsx"
    r2_rows = [{"file": "BBB.json", "subtipo": "factura", "proceso": "p2", "s3_path": "s3://x/BBB"}]
    excel_report.build_maestro_workbook(
        out, group_csv, resumen={"Ambiente": "qa"}, r2_rows=r2_rows,
    )

    wb = load_workbook(out)
    assert "R2" in wb.sheetnames
    ws_r2 = wb["R2"]
    rows = list(ws_r2.iter_rows(values_only=True))
    assert rows[0][0] == "Nombre de los JSONs"
    assert rows[1] == ("BBB.json", "factura", "p2", "s3://x/BBB")


def test_build_maestro_workbook_sin_r2_no_agrega_hoja(tmp_path):
    group_csv = tmp_path / "grupo.csv"
    _write_group_csv(group_csv)

    out = tmp_path / "maestro.xlsx"
    excel_report.build_maestro_workbook(out, group_csv, resumen={"Ambiente": "qa"})

    wb = load_workbook(out)
    assert "R2" not in wb.sheetnames
